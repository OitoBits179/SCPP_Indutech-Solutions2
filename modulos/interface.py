
import os
from datetime import datetime

import customtkinter as ctk
from tkinter import ttk, messagebox, filedialog

from modulos.persistencia import carregar_dados, salvar_dados


# ==========================================
# CONFIGURAÇÕES
# ==========================================

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

SENHA_RESPONSAVEL = "1234"

COR_PRINCIPAL = "#17365D"
COR_SECUNDARIA = "#244A73"
COR_FUNDO = "#F2F5F9"
COR_BRANCA = "#FFFFFF"
COR_TEXTO = "#243447"
COR_VERDE = "#218838"
COR_VERMELHA = "#C0392B"

# Pares de cores (claro, escuro) usados pelos widgets do CustomTkinter.
UI_FUNDO = ("#F2F5F9", "#17191D")
UI_CARTAO = ("#FFFFFF", "#24272D")
UI_TEXTO = ("#243447", "#E7ECF2")
UI_TITULO = ("#17365D", "#EAF2FF")
UI_ENTRADA = ("#FFFFFF", "#30343B")
UI_BORDA = ("#D5DDE7", "#414751")
UI_HOVER = ("#244A73", "#3D6594")
UI_VERDE_HOVER = ("#19692C", "#267D3B")
UI_CINZA = ("#808080", "#666A70")
UI_CINZA_HOVER = ("#666666", "#777C84")
UI_VERMELHO_HOVER = ("#922B21", "#A93226")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CAMINHO_LOGO_LOGIN = os.path.join(BASE_DIR, "assets", "logo_login.png")
CAMINHO_LOGO_CLARO = os.path.join(BASE_DIR, "assets", "logo_claro.png")
CAMINHO_LOGO_ESCURO = os.path.join(BASE_DIR, "assets", "logo_escuro.png")


# ==========================================
# VARIÁVEIS GLOBAIS
# ==========================================

perfil_atual = None
pendencias = []

janela = None
area_conteudo = None
titulo_tela = None
resumo_label = None
perfil_label = None
btn_resolver = None

logo_menu_img = None
logo_login_img = None
logo_menu_label = None
tema_atual = "Claro"
janela_login_ativa = None


# ==========================================
# LOGOS E ESTILOS
# ==========================================
def carregar_imagem(caminho):
    """Abre uma imagem sem manter o arquivo bloqueado."""
    try:
        from PIL import Image

        if not os.path.isfile(caminho):
            print(f"Imagem não encontrada: {caminho}")
            return None
        with Image.open(caminho) as arquivo:
            return arquivo.convert("RGBA")
    except Exception as erro:
        print(f"Não foi possível carregar a imagem {caminho}: {erro}")
        return None


def carregar_logo_menu(tamanho):
    """Cria uma imagem que alterna automaticamente conforme o tema."""
    clara = carregar_imagem(CAMINHO_LOGO_CLARO)
    escura = carregar_imagem(CAMINHO_LOGO_ESCURO)
    if clara is None and escura is None:
        return None
    return ctk.CTkImage(
        light_image=clara or escura,
        dark_image=escura or clara,
        size=tamanho
    )


def carregar_logo_login(tamanho):
    """Carrega a marca simplificada, exclusiva da tela de login."""
    imagem = carregar_imagem(CAMINHO_LOGO_LOGIN)
    if imagem is None:
        return None
    return ctk.CTkImage(light_image=imagem, dark_image=imagem, size=tamanho)


def aplicar_estilo_tabelas():
    """Atualiza as tabelas ttk, que não seguem automaticamente o tema do CTk."""
    try:
        estilo = ttk.Style(janela)
        modo = ctk.get_appearance_mode().lower()
        escuro = modo == "dark"
        fundo = "#24272D" if escuro else "#FFFFFF"
        texto = "#E7ECF2" if escuro else "#243447"
        cabecalho = "#30343B" if escuro else "#E7EDF5"
        selecionado = "#244A73" if not escuro else "#315B87"
        borda = "#414751" if escuro else "#D5DDE7"
        estilo.theme_use("clam")
        estilo.configure(
            "Treeview", background=fundo, foreground=texto,
            fieldbackground=fundo, bordercolor=borda,
            rowheight=30, font=("Arial", 10)
        )
        estilo.configure(
            "Treeview.Heading", background=cabecalho, foreground=texto,
            relief="flat", bordercolor=borda, font=("Arial", 10, "bold")
        )
        estilo.map(
            "Treeview", background=[("selected", selecionado)],
            foreground=[("selected", "#FFFFFF")]
        )
        estilo.map("Treeview.Heading", background=[("active", selecionado)],
                   foreground=[("active", "#FFFFFF" if escuro else texto)])
        estilo.configure("Vertical.TScrollbar", background=cabecalho, troughcolor=fundo)
        estilo.configure("Horizontal.TScrollbar", background=cabecalho, troughcolor=fundo)
    except Exception as erro:
        print(f"Não foi possível atualizar o estilo das tabelas: {erro}")


def aplicar_tema(escolha):
    """Aplica o tema escolhido e atualiza os componentes ttk."""
    global tema_atual
    tema_atual = escolha if escolha in ("Claro", "Escuro", "Sistema") else "Claro"
    modo = {"Claro": "light", "Escuro": "dark", "Sistema": "system"}[tema_atual]
    ctk.set_appearance_mode(modo)
    if janela is not None:
        janela.after(80, aplicar_estilo_tabelas)


# ==========================================
# INICIALIZAÇÃO
# ==========================================

def iniciar():
    global janela, area_conteudo, titulo_tela
    global resumo_label, perfil_label, btn_resolver
    global pendencias, logo_menu_img, logo_menu_label

    dados = carregar_dados()

    # Mantém somente registros no formato esperado.
    if isinstance(dados, list):
        pendencias = [p for p in dados if isinstance(p, dict)]
    else:
        pendencias = []

    janela = ctk.CTk()
    aplicar_tema("Claro")
    janela.title("SCPP - Sistema de Controle de Pendências")
    janela.geometry("1100x680")
    janela.minsize(850, 550)
    janela.configure(fg_color=UI_FUNDO)

    # Menu lateral
    menu = ctk.CTkFrame(
        janela,
        width=220,
        corner_radius=0,
        fg_color=COR_PRINCIPAL
    )
    menu.pack(side="left", fill="y")
    menu.pack_propagate(False)

    # Logo em destaque
    logo_menu_img = carregar_logo_menu((185, 145))

    if logo_menu_img:
        logo_menu_label = ctk.CTkLabel(
            menu, text="", image=logo_menu_img, fg_color=COR_PRINCIPAL
        )
        logo_menu_label.pack(pady=(20, 15))
    else:
        ctk.CTkLabel(
            menu,
            text="SCPP",
            font=("Arial", 32, "bold"),
            text_color=COR_BRANCA
        ).pack(pady=(35, 35))

    criar_botao_menu(menu, "Início", mostrar_inicio)
    criar_botao_menu(menu, "Cadastrar pendência", abrir_cadastro)
    criar_botao_menu(menu, "Consultar pendências", abrir_consulta)

    btn_resolver = ctk.CTkButton(
        menu,
        text="Resolver pendência",
        command=abrir_resolver,
        height=38,
        corner_radius=8,
        fg_color=UI_CINZA,
        hover_color=UI_CINZA_HOVER
    )
    btn_resolver.pack(fill="x", padx=12, pady=5)

    criar_botao_menu(menu, "Exportar PDF", exportar_pdf)
    criar_botao_menu(menu, "Exportar Excel", exportar_excel)

    ctk.CTkLabel(
        menu,
        text="Tema da interface",
        text_color=COR_BRANCA,
        font=("Arial", 12)
    ).pack(pady=(20, 5))

    tema_menu = ctk.CTkOptionMenu(
        menu,
        values=["Claro", "Escuro", "Sistema"],
        command=mudar_tema,
        width=170
    )
    tema_menu.pack(padx=12, pady=5)
    tema_menu.set("Claro")

    ctk.CTkButton(
        menu,
        text="Sair / Trocar usuário",
        command=sair_trocar_usuario,
        height=38,
        corner_radius=8,
        fg_color=COR_VERMELHA,
        hover_color=UI_VERMELHO_HOVER
    ).pack(side="bottom", fill="x", padx=12, pady=18)

    # Área principal
    principal = ctk.CTkFrame(
        janela,
        fg_color=UI_FUNDO,
        corner_radius=0
    )
    principal.pack(
        side="right",
        fill="both",
        expand=True,
        padx=20,
        pady=20
    )

    cabecalho = ctk.CTkFrame(
        principal,
        fg_color=UI_CARTAO,
        corner_radius=12
    )
    cabecalho.pack(fill="x", pady=(0, 15))

    titulo_tela = ctk.CTkLabel(
        cabecalho,
        text="Sistema de Controle de Pendências",
        font=("Arial", 23, "bold"),
        text_color=UI_TITULO
    )
    titulo_tela.pack(anchor="w", padx=20, pady=(15, 5))

    perfil_label = ctk.CTkLabel(
        cabecalho,
        text="Perfil: não selecionado",
        font=("Arial", 13),
        text_color=UI_TEXTO
    )
    perfil_label.pack(anchor="w", padx=20, pady=(0, 5))

    resumo_label = ctk.CTkLabel(
        cabecalho,
        text="Total: 0 | Pendentes: 0 | Resolvidas: 0",
        font=("Arial", 13, "bold"),
        text_color=UI_TEXTO
    )
    resumo_label.pack(anchor="w", padx=20, pady=(0, 15))

    area_conteudo = ctk.CTkFrame(
        principal,
        fg_color=UI_FUNDO,
        corner_radius=0
    )
    area_conteudo.pack(fill="both", expand=True)

    atualizar_resumo()
    mostrar_inicio()

    # Esconde a janela principal até o usuário entrar.
    janela.withdraw()
    janela.after(200, abrir_login)
    janela.mainloop()


def criar_botao_menu(menu, texto, comando):
    ctk.CTkButton(
        menu,
        text=texto,
        command=comando,
        height=38,
        corner_radius=8,
        fg_color=COR_SECUNDARIA,
        hover_color=UI_HOVER
    ).pack(fill="x", padx=12, pady=5)


# ==========================================
# TEMA E ATUALIZAÇÃO
# ==========================================

def mudar_tema(escolha):
    aplicar_tema(escolha)


def salvar():
    """Salva as pendências e retorna True somente se não houver falha."""
    try:
        resultado = salvar_dados(pendencias)

        if resultado is not True:
            messagebox.showerror(
                "Erro ao salvar",
                "Não foi possível salvar os dados. "
                "Verifique o arquivo e tente novamente."
            )
            return False

        return True

    except Exception as erro:
        messagebox.showerror(
            "Erro ao salvar",
            f"Ocorreu um erro ao salvar os dados:\n{erro}"
        )
        return False


def atualizar_resumo():
    if resumo_label is None:
        return

    total = len(pendencias)
    resolvidas = sum(
        1 for p in pendencias
        if p.get("status") == "Resolvida"
    )
    pendentes_count = total - resolvidas

    resumo_label.configure(
        text=(
            f"Total: {total} | "
            f"Pendentes: {pendentes_count} | "
            f"Resolvidas: {resolvidas}"
        )
    )


def limpar_conteudo():
    if area_conteudo is None:
        return

    for widget in area_conteudo.winfo_children():
        widget.destroy()


def definir_titulo(texto):
    if titulo_tela is not None:
        titulo_tela.configure(text=texto)


def mostrar_inicio():
    limpar_conteudo()
    definir_titulo("Início")

    ctk.CTkLabel(
        area_conteudo,
        text="Seja Bem-Vindo",
        font=("Arial", 27, "bold"),
        text_color=UI_TITULO
    ).pack(anchor="w", padx=15, pady=(20, 10))

    ctk.CTkLabel(
        area_conteudo,
        text=(
            "Sistema de Controle de Pendências de Produção.\n\n"
            "Utilize o menu lateral para cadastrar, consultar e acompanhar "
            "as pendências registradas.\n\n"
            "O perfil Responsável possui acesso à resolução das pendências."
        ),
        font=("Arial", 15),
        justify="left",
        text_color=UI_TEXTO,
        wraplength=700
    ).pack(anchor="w", padx=15, pady=10)

    atualizar_resumo()


# ==========================================
# CADASTRO DE PENDÊNCIAS
# ==========================================

def gerar_id():
    maior = 0

    for pendencia in pendencias:
        identificador = str(pendencia.get("id", ""))

        if identificador.startswith("PEN-"):
            try:
                numero = int(identificador.split("-")[1])
                maior = max(maior, numero)
            except (ValueError, IndexError):
                continue

    return f"PEN-{maior + 1:03d}"


def abrir_cadastro():
    limpar_conteudo()
    definir_titulo("Cadastrar pendência")

    ctk.CTkLabel(
        area_conteudo,
        text="Cadastro de nova pendência",
        font=("Arial", 21, "bold"),
        text_color=UI_TITULO
    ).pack(anchor="w", padx=15, pady=(10, 20))

    formulario = ctk.CTkFrame(
        area_conteudo,
        fg_color=UI_CARTAO,
        corner_radius=12
    )
    formulario.pack(fill="x", padx=15, pady=5)

    ctk.CTkLabel(
        formulario,
        text="Produto:",
        font=("Arial", 14, "bold")
    ).pack(anchor="w", padx=20, pady=(15, 5))

    entrada_produto = ctk.CTkEntry(
        formulario,
        placeholder_text="Digite o nome do produto",
        height=38
    )
    entrada_produto.pack(fill="x", padx=20, pady=(0, 12))

    ctk.CTkLabel(
        formulario,
        text="Operador responsável:",
        font=("Arial", 14, "bold")
    ).pack(anchor="w", padx=20, pady=(5, 5))

    entrada_operador = ctk.CTkEntry(
        formulario,
        placeholder_text="Digite o nome do operador",
        height=38
    )
    entrada_operador.pack(fill="x", padx=20, pady=(0, 12))

    ctk.CTkLabel(
        formulario,
        text="Natureza da pendência:",
        font=("Arial", 14, "bold")
    ).pack(anchor="w", padx=20, pady=(5, 5))

    entrada_natureza = ctk.CTkEntry(
        formulario,
        placeholder_text="Descreva a pendência",
        height=38
    )
    entrada_natureza.pack(fill="x", padx=20, pady=(0, 20))

    def registrar():
        produto = entrada_produto.get().strip()
        operador = entrada_operador.get().strip()
        natureza = entrada_natureza.get().strip()

        if not produto or not operador or not natureza:
            messagebox.showwarning(
                "Campos obrigatórios",
                "Preencha todos os campos antes de registrar."
            )
            return

        agora = datetime.now()

        nova_pendencia = {
            "id": gerar_id(),
            "produto": produto,
            "operador": operador,
            "natureza": natureza,
            "data": agora.strftime("%d/%m/%Y"),
            "hora": agora.strftime("%H:%M"),
            "status": "Pendente",
            "data_resolucao": ""
        }

        pendencias.append(nova_pendencia)

        if not salvar():
            pendencias.pop()
            return

        atualizar_resumo()

        messagebox.showinfo(
            "Sucesso",
            f"Pendência {nova_pendencia['id']} cadastrada com sucesso!"
        )

        abrir_cadastro()

    ctk.CTkButton(
        formulario,
        text="Registrar pendência",
        command=registrar,
        height=42,
        fg_color=COR_PRINCIPAL,
        hover_color=UI_HOVER
    ).pack(fill="x", padx=20, pady=(0, 20))


# ==========================================
# CONSULTA DE PENDÊNCIAS
# ==========================================

def abrir_consulta():
    limpar_conteudo()
    definir_titulo("Consultar pendências")

    ctk.CTkLabel(
        area_conteudo,
        text="Pesquisar por ID, produto, operador, natureza ou status",
        font=("Arial", 13),
        text_color=UI_TEXTO
    ).pack(anchor="w", padx=10, pady=(5, 5))

    busca = ctk.CTkEntry(
        area_conteudo,
        placeholder_text="Digite o que deseja pesquisar...",
        height=38
    )
    busca.pack(fill="x", padx=10, pady=(0, 12))

    quadro_tabela = ctk.CTkFrame(
        area_conteudo,
        fg_color=UI_CARTAO
    )
    quadro_tabela.pack(fill="both", expand=True, padx=10, pady=5)

    colunas = (
        "id", "produto", "operador",
        "natureza", "data", "hora", "status"
    )

    aplicar_estilo_tabelas()
    tabela = ttk.Treeview(
        quadro_tabela,
        columns=colunas,
        show="headings"
    )

    titulos = {
        "id": "ID",
        "produto": "Produto",
        "operador": "Operador",
        "natureza": "Natureza",
        "data": "Data",
        "hora": "Hora",
        "status": "Status"
    }

    larguras = {
        "id": 85,
        "produto": 130,
        "operador": 130,
        "natureza": 190,
        "data": 90,
        "hora": 70,
        "status": 95
    }

    for coluna in colunas:
        tabela.heading(coluna, text=titulos[coluna])
        tabela.column(
            coluna,
            width=larguras[coluna],
            minwidth=60,
            anchor="center"
        )

    barra_vertical = ttk.Scrollbar(
        quadro_tabela,
        orient="vertical",
        command=tabela.yview
    )
    barra_horizontal = ttk.Scrollbar(
        quadro_tabela,
        orient="horizontal",
        command=tabela.xview
    )

    tabela.configure(
        yscrollcommand=barra_vertical.set,
        xscrollcommand=barra_horizontal.set
    )

    tabela.grid(row=0, column=0, sticky="nsew")
    barra_vertical.grid(row=0, column=1, sticky="ns")
    barra_horizontal.grid(row=1, column=0, sticky="ew")

    quadro_tabela.grid_rowconfigure(0, weight=1)
    quadro_tabela.grid_columnconfigure(0, weight=1)

    def atualizar_tabela(_evento=None):
        tabela.delete(*tabela.get_children())
        termo = busca.get().strip().lower()

        for p in pendencias:
            valores = (
                p.get("id", ""),
                p.get("produto", ""),
                p.get("operador", ""),
                p.get("natureza", ""),
                p.get("data", ""),
                p.get("hora", ""),
                p.get("status", "")
            )

            if termo and not any(
                termo in str(valor).lower()
                for valor in valores
            ):
                continue

            tabela.insert("", "end", values=valores)

    busca.bind("<KeyRelease>", atualizar_tabela)
    atualizar_tabela()


# ==========================================
# RESOLUÇÃO DE PENDÊNCIAS
# ==========================================

def abrir_resolver():
    if perfil_atual != "Responsável":
        messagebox.showwarning(
            "Acesso restrito",
            "Somente o perfil Responsável pode resolver pendências."
        )
        return

    limpar_conteudo()
    definir_titulo("Resolver pendência")

    ctk.CTkLabel(
        area_conteudo,
        text="Selecione uma pendência pendente e clique em Resolver.",
        font=("Arial", 13),
        text_color=UI_TEXTO
    ).pack(anchor="w", padx=10, pady=(5, 10))

    quadro_tabela = ctk.CTkFrame(
        area_conteudo,
        fg_color=UI_CARTAO
    )
    quadro_tabela.pack(fill="both", expand=True, padx=10, pady=5)

    colunas = ("id", "produto", "operador", "natureza", "data", "status")

    aplicar_estilo_tabelas()
    tabela = ttk.Treeview(
        quadro_tabela,
        columns=colunas,
        show="headings",
        height=12
    )

    titulos = {
        "id": "ID",
        "produto": "Produto",
        "operador": "Operador",
        "natureza": "Natureza",
        "data": "Data",
        "status": "Status"
    }

    for coluna in colunas:
        tabela.heading(coluna, text=titulos[coluna])
        tabela.column(
            coluna,
            width=130,
            minwidth=70,
            anchor="center"
        )

    barra = ttk.Scrollbar(
        quadro_tabela,
        orient="vertical",
        command=tabela.yview
    )
    tabela.configure(yscrollcommand=barra.set)

    tabela.grid(row=0, column=0, sticky="nsew")
    barra.grid(row=0, column=1, sticky="ns")

    quadro_tabela.grid_rowconfigure(0, weight=1)
    quadro_tabela.grid_columnconfigure(0, weight=1)

    # Usa o índice interno da lista como identificador da linha.
    # Assim, registros antigos sem ID não causam conflito na tabela.
    for indice, p in enumerate(pendencias):
        if p.get("status") != "Resolvida":
            tabela.insert(
                "",
                "end",
                iid=str(indice),
                values=(
                    p.get("id", "Sem ID"),
                    p.get("produto", ""),
                    p.get("operador", ""),
                    p.get("natureza", ""),
                    p.get("data", ""),
                    p.get("status", "Pendente")
                )
            )

    def resolver_selecionada():
        if perfil_atual != "Responsável":
            messagebox.showerror(
                "Acesso negado",
                "Seu perfil não tem permissão para resolver pendências."
            )
            return

        selecionados = tabela.selection()

        if not selecionados:
            messagebox.showwarning(
                "Selecione uma pendência",
                "Selecione uma pendência na tabela."
            )
            return

        try:
            indice = int(selecionados[0])
            p = pendencias[indice]
        except (ValueError, IndexError):
            messagebox.showerror(
                "Erro",
                "Não foi possível localizar a pendência selecionada."
            )
            return

        if p.get("status") == "Resolvida":
            messagebox.showinfo(
                "Informação",
                "Essa pendência já foi resolvida."
            )
            return

        status_anterior = p.get("status", "Pendente")
        data_anterior = p.get("data_resolucao", "")

        p["status"] = "Resolvida"
        p["data_resolucao"] = datetime.now().strftime("%d/%m/%Y %H:%M")

        if not salvar():
            p["status"] = status_anterior
            p["data_resolucao"] = data_anterior
            return

        atualizar_resumo()

        messagebox.showinfo(
            "Sucesso",
            f"Pendência {p.get('id', 'selecionada')} resolvida!"
        )

        abrir_resolver()

    ctk.CTkButton(
        area_conteudo,
        text="Resolver pendência selecionada",
        command=resolver_selecionada,
        height=42,
        fg_color=COR_VERDE,
        hover_color=UI_VERDE_HOVER
    ).pack(fill="x", padx=10, pady=10)


# ==========================================
# EXPORTAÇÃO PDF
# ==========================================

def exportar_pdf():
    if not pendencias:
        messagebox.showinfo(
            "Sem dados",
            "Não existem pendências para exportar."
        )
        return

    caminho = filedialog.asksaveasfilename(
        title="Salvar relatório PDF",
        defaultextension=".pdf",
        filetypes=[("Arquivo PDF", "*.pdf")],
        initialfile="relatorio_pendencias.pdf"
    )

    if not caminho:
        return

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.pdfgen import canvas

        largura, altura = landscape(A4)
        pdf = canvas.Canvas(caminho, pagesize=(largura, altura))

        colunas = [
            ("ID", 45),
            ("Produto", 105),
            ("Operador", 100),
            ("Natureza", 205),
            ("Data", 65),
            ("Hora", 45),
            ("Status", 75),
            ("Resolução", 100)
        ]

        def desenhar_cabecalho(numero_pagina):
            pdf.setFillColor(colors.HexColor(COR_PRINCIPAL))
            pdf.rect(0, altura - 65, largura, 65, fill=1, stroke=0)

            pdf.setFillColor(colors.white)
            pdf.setFont("Helvetica-Bold", 16)
            pdf.drawString(
                30,
                altura - 30,
                "Relatório de Pendências - SCPP"
            )

            pdf.setFont("Helvetica", 9)
            pdf.drawString(
                30,
                altura - 48,
                f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            )

            pdf.setFillColor(colors.black)
            pdf.setFont("Helvetica-Bold", 8)

            x = 25
            y = altura - 88

            for titulo, tamanho in colunas:
                pdf.drawString(x, y, titulo)
                x += tamanho

            pdf.setStrokeColor(colors.grey)
            pdf.line(25, y - 5, largura - 25, y - 5)

            pdf.setFont("Helvetica", 7)
            pdf.drawRightString(
                largura - 25,
                18,
                f"Página {numero_pagina}"
            )

        pagina = 1
        desenhar_cabecalho(pagina)
        y = altura - 108

        for p in pendencias:
            if y < 35:
                pdf.showPage()
                pagina += 1
                desenhar_cabecalho(pagina)
                y = altura - 108

            valores = [
                p.get("id", ""),
                p.get("produto", ""),
                p.get("operador", ""),
                p.get("natureza", ""),
                p.get("data", ""),
                p.get("hora", ""),
                p.get("status", ""),
                p.get("data_resolucao", "")
            ]

            x = 25

            for valor, (_, tamanho) in zip(valores, colunas):
                texto = str(valor).replace("\n", " ")
                limite = max(5, int(tamanho / 4.2))

                if len(texto) > limite:
                    texto = texto[:limite - 3] + "..."

                pdf.drawString(x, y, texto)
                x += tamanho

            y -= 18

        pdf.save()

        messagebox.showinfo(
            "Exportação concluída",
            "O relatório PDF foi salvo com sucesso."
        )

    except ImportError:
        messagebox.showerror(
            "Biblioteca ausente",
            "Instale a biblioteca necessária com:\n"
            "pip install reportlab"
        )

    except Exception as erro:
        messagebox.showerror(
            "Erro na exportação",
            f"Não foi possível gerar o PDF:\n{erro}"
        )


# ==========================================
# EXPORTAÇÃO EXCEL
# ==========================================

def exportar_excel():
    if not pendencias:
        messagebox.showinfo(
            "Sem dados",
            "Não existem pendências para exportar."
        )
        return

    caminho = filedialog.asksaveasfilename(
        title="Salvar relatório Excel",
        defaultextension=".xlsx",
        filetypes=[("Arquivo Excel", "*.xlsx")],
        initialfile="relatorio_pendencias.xlsx"
    )

    if not caminho:
        return

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        livro = Workbook()
        planilha = livro.active
        planilha.title = "Pendências"

        cabecalhos = [
            "ID",
            "Produto",
            "Operador",
            "Natureza",
            "Data",
            "Hora",
            "Status",
            "Data da resolução"
        ]

        planilha.append(cabecalhos)

        for p in pendencias:
            planilha.append([
                p.get("id", ""),
                p.get("produto", ""),
                p.get("operador", ""),
                p.get("natureza", ""),
                p.get("data", ""),
                p.get("hora", ""),
                p.get("status", ""),
                p.get("data_resolucao", "")
            ])

        preenchimento = PatternFill(
            fill_type="solid",
            fgColor="17365D"
        )

        for celula in planilha[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = preenchimento
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        larguras = [14, 24, 24, 40, 15, 12, 18, 24]

        for indice, largura in enumerate(larguras, start=1):
            planilha.column_dimensions[
                get_column_letter(indice)
            ].width = largura

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions

        livro.save(caminho)

        messagebox.showinfo(
            "Exportação concluída",
            "O arquivo Excel foi salvo com sucesso."
        )

    except ImportError:
        messagebox.showerror(
            "Biblioteca ausente",
            "Instale a biblioteca necessária com:\n"
            "pip install openpyxl"
        )

    except Exception as erro:
        messagebox.showerror(
            "Erro na exportação",
            f"Não foi possível gerar o Excel:\n{erro}"
        )


# ==========================================
# PERMISSÕES
# ==========================================

def aplicar_permissoes():
    if btn_resolver is None:
        return

    if perfil_atual == "Responsável":
        btn_resolver.configure(
            state="normal",
            fg_color=COR_VERDE,
            hover_color=UI_VERDE_HOVER
        )
    else:
        btn_resolver.configure(
            state="disabled",
            fg_color=UI_CINZA,
            hover_color=UI_CINZA_HOVER
        )


def atualizar_perfil():
    if perfil_label is not None:
        nome = perfil_atual if perfil_atual else "não selecionado"
        perfil_label.configure(text=f"Perfil: {nome}")

    aplicar_permissoes()


# ==========================================
# LOGIN E TROCA DE USUÁRIO
# ==========================================

def entrar_funcionario(janela_login):
    global perfil_atual

    perfil_atual = "Funcionário"
    atualizar_perfil()

    janela_login.destroy()
    aplicar_tema(tema_atual)
    janela.deiconify()
    mostrar_inicio()


def abrir_tela_senha(janela_login):
    janela_login.withdraw()

    janela_senha = ctk.CTkToplevel(janela_login)
    janela_senha.configure(fg_color="#FFFFFF")
    janela_senha.title("Acesso do Responsável")
    janela_senha.geometry("350x250")
    janela_senha.resizable(False, False)
    janela_senha.transient(janela_login)
    janela_senha.grab_set()

    ctk.CTkLabel(
        janela_senha,
        text="Acesso do Responsável",
        font=("Arial", 19, "bold"),
        text_color=UI_TITULO
    ).pack(pady=(25, 15))

    ctk.CTkLabel(
        janela_senha,
        text="Digite a senha para continuar:"
    ).pack(pady=5)

    entrada_senha = ctk.CTkEntry(
        janela_senha,
        placeholder_text="Senha",
        show="*",
        height=38
    )
    entrada_senha.pack(fill="x", padx=25, pady=10)
    entrada_senha.focus_set()

    def confirmar():
        global perfil_atual

        if entrada_senha.get() != SENHA_RESPONSAVEL:
            messagebox.showerror(
                "Senha incorreta",
                "A senha digitada está incorreta.",
                parent=janela_senha
            )
            entrada_senha.delete(0, "end")
            entrada_senha.focus_set()
            return

        perfil_atual = "Responsável"
        atualizar_perfil()

        janela_senha.destroy()
        janela_login.destroy()
        aplicar_tema(tema_atual)
        janela.deiconify()
        mostrar_inicio()

    def voltar():
        janela_senha.destroy()
        janela_login.deiconify()
        janela_login.grab_set()

    ctk.CTkButton(
        janela_senha,
        text="Entrar",
        command=confirmar,
        height=38,
        fg_color=COR_PRINCIPAL
    ).pack(fill="x", padx=25, pady=(5, 8))

    ctk.CTkButton(
        janela_senha,
        text="Voltar",
        command=voltar,
        height=34,
        fg_color=UI_CINZA,
        hover_color=UI_CINZA_HOVER
    ).pack(fill="x", padx=25, pady=(0, 15))

    janela_senha.bind("<Return>", lambda _evento: confirmar())
    janela_senha.protocol("WM_DELETE_WINDOW", voltar)


def abrir_login():
    global logo_login_img, janela_login_ativa

    # A tela de login permanece clara, mesmo que o sistema esteja no tema escuro.
    ctk.set_appearance_mode("light")
    janela.withdraw()

    janela_login = ctk.CTkToplevel(janela)
    janela_login_ativa = janela_login
    janela_login.configure(fg_color="#FFFFFF")
    janela_login.title("Login - SCPP")
    janela_login.geometry("460x500")
    janela_login.resizable(False, False)
    janela_login.transient(janela)
    janela_login.grab_set()
    janela_login.protocol("WM_DELETE_WINDOW", janela.destroy)

    ctk.CTkLabel(
        janela_login,
        text="Bem-vindo ao SCPP",
        font=("Arial", 24, "bold"),
        text_color=UI_TITULO
    ).pack(pady=(18, 8))

    # Logo grande na tela de login
    logo_login_img = carregar_logo_login((210, 210))

    if logo_login_img:
        ctk.CTkLabel(
            janela_login,
            text="",
            image=logo_login_img
        ).pack(pady=(0, 10))
    else:
        ctk.CTkLabel(
            janela_login,
            text="SCPP",
            font=("Arial", 38, "bold"),
            text_color=UI_TITULO
        ).pack(pady=(10, 25))

    ctk.CTkLabel(
        janela_login,
        text="Selecione seu perfil para continuar:",
        font=("Arial", 14)
    ).pack(pady=(5, 12))

    ctk.CTkButton(
        janela_login,
        text="FUNCIONÁRIO",
        command=lambda: entrar_funcionario(janela_login),
        height=42,
        fg_color=COR_PRINCIPAL,
        hover_color=UI_HOVER
    ).pack(fill="x", padx=45, pady=6)

    ctk.CTkButton(
        janela_login,
        text="RESPONSÁVEL",
        command=lambda: abrir_tela_senha(janela_login),
        height=42,
        fg_color=COR_VERDE,
        hover_color=UI_VERDE_HOVER
    ).pack(fill="x", padx=45, pady=6)


def sair_trocar_usuario():
    global perfil_atual

    resposta = messagebox.askyesno(
        "Trocar usuário",
        "Deseja sair do perfil atual e voltar à tela de login?",
        parent=janela
    )

    if not resposta:
        return

    perfil_atual = None
    atualizar_perfil()
    abrir_login()


# ==========================================
# EXECUÇÃO
# ==========================================

if __name__ == "__main__":
    iniciar()
